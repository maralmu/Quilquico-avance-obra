"""Pruebas de regresión contra el código entregado y escenarios inválidos."""
import ast
import io
from pathlib import Path
import unittest
import pandas as pd
from pandas.testing import assert_frame_equal, assert_series_equal
from openpyxl import load_workbook
import calculos as c
from graficos import comparar, grafico_desviacion, grafico_partidas_encargado
from reportes import generar_reporte_excel

ROOT = Path(__file__).resolve().parents[1]


class ProyectoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = c.cargar_datos(ROOT / 'datos/Caso_2_Quilquico_Base_Completa.csv')

    def test_regresion_original(self):
        codigo = ast.parse((ROOT / 'referencias/main_tarea_final_corregido.py').read_text(encoding='utf-8-sig'))
        nombres = ['calcular_avance_semanal', 'detectar_atrasos', 'avance_por_especialidad', 'avance_por_encargado', 'partidas_por_encargado']
        modulo = ast.Module(body=[n for n in codigo.body if isinstance(n, ast.FunctionDef) and n.name in nombres], type_ignores=[])
        espacio = {'pd': pd}
        exec(compile(modulo, '<funciones originales>', 'exec'), espacio)
        for nombre in nombres:
            with self.subTest(funcion=nombre):
                actual = getattr(c, nombre)(self.d)
                original = espacio[nombre](self.d)
                if isinstance(actual, pd.Series):
                    assert_series_equal(actual, original)
                else:
                    assert_frame_equal(actual, original)

    def test_resultados_y_corte(self):
        a = c.calcular_avance_semanal(self.d)
        self.assertAlmostEqual(a.loc[8, 'planificado_acumulado'], .87)
        self.assertAlmostEqual(a.loc[8, 'real_acumulado'], .815)
        self.assertAlmostEqual(a.loc[12, 'real_acumulado'], 1)
        self.assertEqual(c.estado(a.loc[12, 'desviacion']), 'SEGÚN PROGRAMA')
        self.assertAlmostEqual(c.resumen_partidas(self.d, 8).real_acumulado.sum(), .815)
        self.assertEqual(c.partidas_por_encargado(self.d).sum(), 20)

    def test_invalidos(self):
        casos = []
        casos.append(self.d.iloc[0:0])
        casos.append(self.d.drop(columns='semana'))
        casos.append(pd.concat([self.d, self.d.iloc[[0]]]))
        casos.append(self.d[self.d.semana != 2])
        casos.append(self.d.iloc[1:])
        for campo, valor in [('semana', 1.5), ('semana', 0), ('peso_partida', 50.), (c.AVANCES[0], -1.), (c.AVANCES[1], 2.), (c.AVANCES[1], float('inf')), ('partida', ' '), ('encargado', None)]:
            d = self.d.copy()
            d[campo] = d[campo].astype(object)
            d.loc[0, campo] = valor
            casos.append(d)
        for i, d in enumerate(casos):
            with self.subTest(caso=i), self.assertRaises(ValueError):
                c.validar_datos(d)
        for contenido in [b'', b'columna\ntexto', b'\xff\xfe']:
            with self.subTest(contenido=contenido), self.assertRaises(ValueError):
                c.cargar_datos(contenido)

    def test_edicion_transaccional(self):
        antes = self.d.copy(deep=True)
        fila = self.d.iloc[0]
        with self.assertRaises(ValueError):
            c.actualizar_registro(self.d, fila.semana, fila.codigo_partida, 1., 1., fila.especialidad, fila.encargado)
        assert_frame_equal(self.d, antes)
        editado = c.actualizar_registro(self.d, fila.semana, fila.codigo_partida, 0., 0., fila.especialidad, fila.encargado)
        self.assertEqual(len(editado), 240)
        assert_frame_equal(self.d, antes)

    def test_excel_y_graficos(self):
        wb = load_workbook(io.BytesIO(generar_reporte_excel(self.d, 8)))
        self.assertEqual(len(wb.worksheets), 6)
        self.assertEqual(len(wb['Avance semanal']._charts), 2)
        self.assertEqual(wb['Avance semanal'].max_row, 10)
        self.assertAlmostEqual(wb['Resumen']['B6'].value, .815)
        self.assertEqual(wb['Avance semanal']['B3'].number_format, '0.00%')
        a = c.calcular_avance_semanal(self.d)
        for fig in [comparar(a, 'Curva S', acumulado=True), comparar(a, 'Semanal', barras=True), grafico_desviacion(a), grafico_partidas_encargado(c.partidas_por_encargado(self.d)), comparar(c.avance_por_especialidad(self.d), 'Especialidades', barras=True)]:
            self.assertTrue(fig.data)
            self.assertIn('data', fig.to_json())

    def test_streamlit(self):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30).run()
        self.assertFalse(at.exception)
        at.sidebar.selectbox[0].select(8).run()
        self.assertIn('81.50%', [m.value for m in at.metric])
        for seccion in ['Avance general', 'Especialidades', 'Encargados', 'Partidas y alertas', 'Reportes']:
            at.sidebar.radio[1].set_value(seccion).run()
            self.assertFalse(at.exception, seccion)
        next(b for b in at.button if b.label == 'Generar reporte Excel').click().run()
        self.assertFalse(at.exception)
        self.assertIn('excel', at.session_state)
        at.sidebar.radio[0].set_value('Cargar CSV').run()
        self.assertFalse(at.exception)
        self.assertTrue(at.info)


if __name__ == '__main__':
    unittest.main(verbosity=2)
