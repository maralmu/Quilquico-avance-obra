"""Reporte de resultados calculados, con gráficos nativos editables en Excel."""
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from calculos import (AVANCES, calcular_avance_semanal, avance_por_especialidad,
                      avance_por_encargado, partidas_por_encargado, resumen_partidas, estado)


def _hoja(libro, nombre, titulo, tabla, porcentajes=()):
    h = libro.create_sheet(nombre)
    h.append([titulo])
    h.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(tabla.columns))
    h.row_dimensions[1].height = 32
    h.cell(1, 1).font = Font(size=16, bold=True, color='17324D')
    etiquetas = {'semana': 'Semana', 'codigo_partida': 'Código partida', 'partida': 'Partida',
        'especialidad': 'Especialidad', 'encargado': 'Encargado', 'peso_partida': 'Peso (%)',
        AVANCES[0]: 'Planificado semanal', AVANCES[1]: 'Real semanal',
        'planificado_acumulado': 'Planificado acumulado', 'real_acumulado': 'Real acumulado',
        'desviacion': 'Desviación', 'cantidad_partidas': 'Cantidad de partidas', 'estado': 'Estado'}
    h.append([etiquetas.get(c, c) for c in tabla.columns])
    for fila in tabla.itertuples(index=False, name=None):
        h.append([v.item() if hasattr(v, 'item') else v for v in fila])
    for celda in h[2]:
        celda.fill = PatternFill('solid', fgColor='17324D')
        celda.font = Font(bold=True, color='FFFFFF')
        celda.alignment = Alignment(wrap_text=True, vertical='center')
    h.row_dimensions[2].height = 36
    for j, c in enumerate(tabla.columns, 1):
        longitud = max([len(str(v)) for v in tabla[c]] + [len(etiquetas.get(c, c))])
        h.column_dimensions[get_column_letter(j)].width = min(55, max(20, longitud + 2))
        for row in h.iter_rows(min_row=3, min_col=j, max_col=j):
            celda = row[0]
            # Evita que nombres aportados por un CSV se interpreten como fórmulas.
            if isinstance(celda.value, str):
                celda.data_type = 's'
                celda.alignment = Alignment(wrap_text=True, vertical='center')
                if len(celda.value) > 53:
                    h.row_dimensions[celda.row].height = 42
            if c in porcentajes:
                celda.number_format = '0.00%'
            if celda.row % 2:
                celda.fill = PatternFill('solid', fgColor='EFF5F8')
    h.freeze_panes = 'B3'
    h.auto_filter.ref = f'A2:{get_column_letter(len(tabla.columns))}{h.max_row}'
    h.sheet_view.showGridLines = False
    return h


def generar_reporte_excel(datos, semana=None, proyecto='Proyecto Quilquico'):
    semana = int(datos.semana.max()) if semana is None else int(semana)
    d = datos[datos.semana <= semana].copy()
    if d.empty:
        raise ValueError('No hay datos en el corte seleccionado.')
    a = calcular_avance_semanal(d)
    final = a.iloc[-1]
    wb = Workbook()
    wb.remove(wb.active)
    resumen = pd.DataFrame({'Indicador': ['Proyecto', 'Semana de corte', 'Planificado acumulado', 'Real acumulado', 'Desviación (fracción)', 'Estado', 'Partidas', 'Especialidades', 'Encargados'],
        'Resultado': [proyecto, semana, final.planificado_acumulado, final.real_acumulado, final.desviacion, estado(final.desviacion), d.codigo_partida.nunique(), d.especialidad.nunique(), d.encargado.nunique()]})
    h = _hoja(wb, 'Resumen', f'{proyecto} · Control de avance', resumen)
    h.column_dimensions['A'].width = 32
    h.column_dimensions['B'].width = 44
    for r in (5, 6, 7):
        h.cell(r, 2).number_format = '0.00%'
    h.cell(14, 1, 'Los avances son contribuciones ponderadas al total de la obra.')
    h.cell(15, 1, 'La diferencia entre porcentajes se interpreta en puntos porcentuales.')
    h.cell(16, 1, 'Reporte de valores calculados; regenere desde la app si cambian los datos.')
    h = _hoja(wb, 'Avance semanal', f'Avance general hasta semana {semana}', a.reset_index(), a.columns)
    for cls, cols, titulo, ancla in [(LineChart, (4, 5), 'Curva S: planificado y real', 'I2'), (BarChart, (2, 3), 'Comparación de avance semanal', 'I23')]:
        chart = cls()
        chart.title = titulo
        chart.y_axis.title = 'Avance (%)'
        chart.y_axis.numFmt = '0%'
        chart.x_axis.title = 'Semana'
        chart.add_data(Reference(h, min_col=cols[0], max_col=cols[1], min_row=2, max_row=h.max_row), titles_from_data=True)
        for serie, nombre in zip(chart.series, ['Planificado', 'Real']):
            serie.tx = SeriesLabel(v=nombre)
        chart.set_categories(Reference(h, min_col=1, min_row=3, max_row=h.max_row))
        chart.width, chart.height = 23, 10
        if isinstance(chart, BarChart):
            chart.type, chart.grouping, chart.overlap = 'col', 'clustered', 0
        h.add_chart(chart, ancla)
    h.conditional_formatting.add(f'F3:F{h.max_row}', ColorScaleRule(start_type='min', start_color='F6C5BF', end_type='max', end_color='C3E6DD'))
    for nombre, grupo, funcion in [('Especialidades', 'especialidad', avance_por_especialidad), ('Encargados', 'encargado', avance_por_encargado)]:
        t = funcion(d)[AVANCES].copy()
        t['desviacion'] = t[AVANCES[1]] - t[AVANCES[0]]
        if grupo == 'encargado':
            t['cantidad_partidas'] = partidas_por_encargado(d)
        t = t.rename(columns={AVANCES[0]: 'planificado_acumulado', AVANCES[1]: 'real_acumulado'})
        _hoja(wb, nombre, f'{nombre}: aporte acumulado al total de obra', t.reset_index(), ['planificado_acumulado', 'real_acumulado', 'desviacion'])
    p = resumen_partidas(d, semana)
    _hoja(wb, 'Partidas', f'Estado acumulado por partida: semana {semana}', p.reset_index(), ['peso_partida', 'planificado_acumulado', 'real_acumulado', 'desviacion'])
    _hoja(wb, 'Base de datos', 'Fuente: datos cargados en la aplicación, hasta el corte seleccionado', d, ['peso_partida', *AVANCES])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
