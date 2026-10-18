# Se selecciono el caso 2 por su cercania con la realidad de mis labores.
# Se adapto un proyecto que estoy realizando a las necesidades solicitadas.

import pandas as pd
import matplotlib.pyplot as plt
import plotly.graph_objects as go
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side



def cargar_datos(archivo):
    datos = pd.read_csv(archivo)
    return datos


def calcular_avance_semanal(datos):
    avance_semanal = datos.groupby("semana")[
        ["avance_planificado_semanal", "avance_real_semanal"]
    ].sum()

    avance_semanal["planificado_acumulado"] = (
        avance_semanal["avance_planificado_semanal"].cumsum()
    )

    avance_semanal["real_acumulado"] = (
        avance_semanal["avance_real_semanal"].cumsum()
    )

    avance_semanal["desviacion"] = (
        avance_semanal["real_acumulado"]
        - avance_semanal["planificado_acumulado"]
    )

    return avance_semanal


def revisar_estado(avance_semanal):
    for semana, fila in avance_semanal.iterrows():

        if fila["desviacion"] < 0:
            estado = "ATRASADA"
        elif fila["desviacion"] > 0:
            estado = "ADELANTADA"
        else:
            estado = "SEGUN PROGRAMA"

        print(
            f"Semana {semana}: Estado de la obra: {estado}, "
            f"Desviacion: {round(fila['desviacion'] * 100, 2)} %"
        )


def detectar_atrasos(datos):
    datos = datos.copy()

    datos["diferencia"] = (
        datos["avance_real_semanal"]
        - datos["avance_planificado_semanal"]
    )

    partidas_atrasadas = datos[datos["diferencia"] < 0].copy()

    return partidas_atrasadas


def avance_por_especialidad(datos):
    avance = datos.groupby("especialidad")[
        ["avance_planificado_semanal", "avance_real_semanal"]
    ].sum()

    avance["planificado_porcentaje"] = (
        avance["avance_planificado_semanal"] * 100
    )

    avance["real_porcentaje"] = (
        avance["avance_real_semanal"] * 100
    )

    return avance


def partidas_por_encargado(datos):
    cantidad = datos.groupby("encargado")["codigo_partida"].nunique()
    return cantidad


def avance_por_encargado(datos):
    avance = datos.groupby("encargado")[
        ["avance_planificado_semanal", "avance_real_semanal"]
    ].sum()

    avance["planificado_porcentaje"] = (
        avance["avance_planificado_semanal"] * 100
    )

    avance["real_porcentaje"] = (
        avance["avance_real_semanal"] * 100
    )

    return avance


def grafico_curva_s(avance_semanal):
    semanas = [0] + list(avance_semanal.index)
    planificado = [0] + list(
        avance_semanal["planificado_acumulado"] * 100
    )
    real = [0] + list(
        avance_semanal["real_acumulado"] * 100
    )

    plt.figure(figsize=(9, 5))

    plt.plot(semanas, planificado, marker="o", label="Planificado")
    plt.plot(semanas, real, marker="o", label="Real")

    plt.xlabel("Semana")
    plt.ylabel("Avance acumulado (%)")
    plt.title("Curva S")
    plt.xticks(semanas)
    plt.legend()
    plt.grid()
    plt.tight_layout()
    plt.show()


def grafico_avance_semanal(avance_semanal):
    semanas = avance_semanal.index
    ancho = 0.35

    plt.figure(figsize=(9, 5))

    plt.bar(
        semanas - ancho / 2,
        avance_semanal["avance_planificado_semanal"] * 100,
        width=ancho,
        label="Planificado"
    )

    plt.bar(
        semanas + ancho / 2,
        avance_semanal["avance_real_semanal"] * 100,
        width=ancho,
        label="Real"
    )

    plt.xlabel("Semana")
    plt.ylabel("Avance semanal (%)")
    plt.title("Avance planificado vs real")
    plt.xticks(semanas)
    plt.legend()
    plt.grid(axis="y")
    plt.tight_layout()
    plt.show()


def grafico_especialidad(datos):
    datos_especialidad = datos.groupby(
        ["semana", "especialidad"]
    )[
        ["avance_planificado_semanal", "avance_real_semanal"]
    ].sum().reset_index()

    for especialidad in datos_especialidad["especialidad"].unique():

        datos_grafico = datos_especialidad[
            datos_especialidad["especialidad"] == especialidad
        ]

        plt.figure(figsize=(9, 5))

        plt.plot(
            datos_grafico["semana"],
            datos_grafico["avance_planificado_semanal"] * 100,
            marker="o",
            label="Planificado"
        )

        plt.plot(
            datos_grafico["semana"],
            datos_grafico["avance_real_semanal"] * 100,
            marker="o",
            label="Real"
        )

        plt.xlabel("Semana")
        plt.ylabel("Avance semanal (%)")
        plt.title(f"Avance planificado vs real - {especialidad}")
        plt.xticks(datos_grafico["semana"])
        plt.legend()
        plt.grid()
        plt.tight_layout()
        plt.show()


def grafico_partidas_encargado(cantidad_partidas):
    plt.figure(figsize=(9, 5))

    cantidad_partidas.plot(kind="bar")

    plt.xlabel("Encargado")
    plt.ylabel("Cantidad de partidas")
    plt.title("Cantidad de partidas por encargado")
    plt.xticks(rotation=45)
    plt.grid(axis="y")
    plt.tight_layout()
    plt.show()


def grafico_encargado_plotly(datos):
    datos_encargado = datos.groupby(
        ["semana", "encargado"]
    )[
        ["avance_planificado_semanal", "avance_real_semanal"]
    ].sum().reset_index()

    datos_encargado = datos_encargado.sort_values(
        ["encargado", "semana"]
    )

    datos_encargado["planificado_acumulado"] = (
        datos_encargado.groupby("encargado")[
            "avance_planificado_semanal"
        ].cumsum() * 100
    )

    datos_encargado["real_acumulado"] = (
        datos_encargado.groupby("encargado")[
            "avance_real_semanal"
        ].cumsum() * 100
    )

    encargados = datos_encargado["encargado"].unique()
    fig = go.Figure()

    for i, encargado in enumerate(encargados):
        datos_grafico = datos_encargado[
            datos_encargado["encargado"] == encargado
        ]

        visible = i == 0

        fig.add_trace(
            go.Scatter(
                x=datos_grafico["semana"],
                y=datos_grafico["planificado_acumulado"],
                mode="lines+markers",
                name="Planificado",
                visible=visible
            )
        )

        fig.add_trace(
            go.Scatter(
                x=datos_grafico["semana"],
                y=datos_grafico["real_acumulado"],
                mode="lines+markers",
                name="Real",
                visible=visible
            )
        )

    botones = []

    for i, encargado in enumerate(encargados):
        visible = [False] * (len(encargados) * 2)
        visible[i * 2] = True
        visible[i * 2 + 1] = True

        botones.append(
            dict(
                label=encargado,
                method="update",
                args=[
                    {"visible": visible},
                    {"title": f"Avance acumulado - {encargado}"}
                ]
            )
        )

    fig.update_layout(
        title=f"Avance acumulado - {encargados[0]}",
        xaxis_title="Semana",
        yaxis_title="Avance acumulado (%)",
        hovermode="x unified",
        updatemenus=[
            dict(
                buttons=botones,
                direction="down",
                showactive=True
            )
        ]
    )

    fig.show()


# Inicio del programa
datos = cargar_datos("caso_2_Quilquico_base_Completa.csv")

print(datos.head())

avance_semanal = calcular_avance_semanal(datos)

print("\nAVANCE SEMANAL")
print(avance_semanal.round(4))

revisar_estado(avance_semanal)

partidas_atrasadas = detectar_atrasos(datos)

print("\nPARTIDAS CON ATRASO")
if partidas_atrasadas.empty:
    print("No se detectaron atrasos.")
else:
    print(
        partidas_atrasadas[
            ["semana", "codigo_partida", "partida", "encargado", "diferencia"]
        ].to_string(index=False)
    )

avance_especialidad = avance_por_especialidad(datos)

print("\nAVANCE POR ESPECIALIDAD")
print(
    avance_especialidad[
        ["planificado_porcentaje", "real_porcentaje"]
    ].round(2)
)

cantidad_partidas = partidas_por_encargado(datos)

print("\nCANTIDAD DE PARTIDAS POR ENCARGADO")
print(cantidad_partidas)

avance_encargado = avance_por_encargado(datos)

print("\nAVANCE POR ENCARGADO")
print(
    avance_encargado[
        ["planificado_porcentaje", "real_porcentaje"]
    ].round(2)
)

print("\n#---------------------------------------")
print("GRAFICO 1: CURVA S")
print("#---------------------------------------")
grafico_curva_s(avance_semanal)

print("\n#---------------------------------------")
print("GRAFICO 2: AVANCE SEMANAL")
print("#---------------------------------------")
grafico_avance_semanal(avance_semanal)

print("\n#---------------------------------------")
print("GRAFICO 3: AVANCE POR ESPECIALIDAD")
print("#---------------------------------------")
grafico_especialidad(datos)

print("\n#---------------------------------------")
print("GRAFICO 4: PARTIDAS POR ENCARGADO")
print("#---------------------------------------")
grafico_partidas_encargado(cantidad_partidas)

print("\n#---------------------------------------")
print("GRAFICO 5: AVANCE POR ENCARGADO - PLOTLY")
print("#---------------------------------------")
grafico_encargado_plotly(datos)

# Creacion del Excel

def generar_reporte_excel(
    datos,
    avance_semanal,
    avance_especialidad,
    cantidad_partidas,
    avance_encargado,
    nombre_archivo
):
    libro = Workbook()

    # Hoja 1: Base de datos

    hoja_base = libro.active
    hoja_base.title = "Base de datos"

    for columna, nombre in enumerate(datos.columns, start=1):
        hoja_base.cell(row=1, column=columna, value=nombre)

    for fila, valores in enumerate(datos.itertuples(index=False), start=2):
        for columna, valor in enumerate(valores, start=1):
            hoja_base.cell(row=fila, column=columna, value=valor)

    # Hoja 2: Avance semanal

    hoja_avance = libro.create_sheet("Avance semanal")

    encabezados = [
        "Semana",
        "Planificado semanal",
        "Real semanal",
        "Planificado acumulado",
        "Real acumulado",
        "Desviacion"
    ]

    hoja_avance.append(encabezados)

    for semana, fila in avance_semanal.iterrows():
        hoja_avance.append(
            [
                semana,
                fila["avance_planificado_semanal"],
                fila["avance_real_semanal"],
                fila["planificado_acumulado"],
                fila["real_acumulado"],
                fila["desviacion"]
            ]
        )

    for fila in hoja_avance.iter_rows(
        min_row=2,
        min_col=2,
        max_col=6
    ):
        for celda in fila:
            celda.number_format = "0.00%"

    # Curva S
    grafico = LineChart()

    grafico.title = "Curva S - Avance planificado vs real"
    grafico.y_axis.title = "Avance acumulado"
    grafico.x_axis.title = "Semana"

    datos_grafico = Reference(
        hoja_avance,
        min_col=4,
        max_col=5,
        min_row=1,
        max_row=hoja_avance.max_row
    )

    categorias = Reference(
        hoja_avance,
        min_col=1,
        min_row=2,
        max_row=hoja_avance.max_row
    )

    grafico.add_data(datos_grafico, titles_from_data=True)
    grafico.set_categories(categorias)

    grafico.height = 10
    grafico.width = 18

    hoja_avance.add_chart(grafico, "H2")

    # Hoja 3: Especialidades

    hoja_especialidad = libro.create_sheet("Especialidades")

    hoja_especialidad.append(
        [
            "Especialidad",
            "Planificado (%)",
            "Real (%)"
        ]
    )

    for especialidad, fila in avance_especialidad.iterrows():
        hoja_especialidad.append(
            [
                especialidad,
                fila["planificado_porcentaje"],
                fila["real_porcentaje"]
            ]
        )

    # Hoja 4: Encargados

    hoja_encargado = libro.create_sheet("Encargados")

    hoja_encargado.append(
        [
            "Encargado",
            "Cantidad de partidas",
            "Planificado (%)",
            "Real (%)"
        ]
    )

    for encargado, fila in avance_encargado.iterrows():
        hoja_encargado.append(
            [
                encargado,
                int(cantidad_partidas[encargado]),
                fila["planificado_porcentaje"],
                fila["real_porcentaje"]
            ]
        )

    color_encabezado = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    borde = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for hoja in libro.worksheets:

        for celda in hoja[1]:
            celda.font = Font(bold=True)
            celda.fill = color_encabezado
            celda.alignment = Alignment(horizontal="center")
            celda.border = borde

        for columna in hoja.columns:
            ancho = 0

            for celda in columna:
                if celda.value is not None:
                    ancho = max(
                        ancho,
                        len(str(celda.value))
                    )

            hoja.column_dimensions[
                columna[0].column_letter
            ].width = ancho + 2

    libro.save(nombre_archivo)

    print(f"Reporte Excel generado: {nombre_archivo}")


# Generar reporte Excel
generar_reporte_excel(
    datos,
    avance_semanal,
    avance_especialidad,
    cantidad_partidas,
    avance_encargado,
    "Reporte_Proyecto_Quilquico.xlsx"
)

print("\n#---------------------------------------")
print("FIN DEL ANALISIS")
print("#---------------------------------------")